from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt

DOC = Path(__file__).resolve().parents[1] / "documents/pcb_thesis_deeppcb_revised.docx"

TEXT = ("论文修改时还应保持数据、代码和文字之间的一致性。数据集章节说明的类别顺序必须与数据配置文件一致，"
        "方法章节给出的层级和步长必须与模型配置一致，实验章节的表格必须能够追溯到对应运行目录。若重新训练后更换了输入尺寸、"
        "随机种子或权重，摘要和结论中的数字也要同步更新。对于没有运行的融合消融和外部数据实验，应继续保留待实测标记，"
        "避免把研究计划误写成已经验证的事实。这样既能增加论文内容的完整性，也能使审阅者根据清单、命令和日志复查每一项结果。")

d = Document(DOC)
targets = [p for p in d.paragraphs if p.text.strip() == "6.3 后续研究方向"]
if len(targets) != 1:
    raise RuntimeError(f"unexpected target count: {len(targets)}")
anchor = targets[0]
p = d.add_paragraph(style="Normal")
p.add_run(TEXT)
anchor._p.addnext(p._p)
p.paragraph_format.first_line_indent = Cm(0.74)
p.paragraph_format.line_spacing = 1.42
p.paragraph_format.space_after = Pt(0)
if not any("正文的篇章一致性" in q.text for q in d.paragraphs):
    extra = d.add_paragraph(style="Normal")
    extra.add_run("正文的篇章一致性还体现在章节之间的相互照应：绪论提出的问题应在方法章节得到对应设计，方法中的变量应在实验章节得到单独对照，实验现象又应反馈到误差分析和展望。按这一关系整理文字，可以减少重复叙述，也能让读者清楚知道每一项实验究竟用于回答哪个研究问题。")
    anchor._p.addnext(extra._p)
    extra.paragraph_format.first_line_indent = Cm(0.74)
    extra.paragraph_format.line_spacing = 1.42
    extra.paragraph_format.space_after = Pt(0)
d.save(DOC)
print(DOC)
