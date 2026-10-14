"""Write the measured 10-epoch CPU results into the editable thesis tables."""
from pathlib import Path
from docx import Document

DOC = Path(__file__).resolve().parents[1] / "documents/pcb_thesis_deeppcb_revised.docx"

BASE = {
    "m50": 25.5285, "m5095": 11.5029, "p": 23.7361, "r": 36.5759,
    "ap": [16.3068, 9.4052, 13.3152, 9.4095, 70.6151, 34.1191],
    "rec": [25.9484, 20.7113, 25.5973, 6.6253, 78.2328, 62.3404],
    "params": 2.58, "flops": 6.4, "latency": 17.5,
}
P2 = {
    "m50": 20.4037, "m5095": 9.5723, "p": 20.7381, "r": 36.5162,
    "ap": [16.1784, 6.6183, 15.2723, 8.5665, 36.4477, 39.3392],
    "rec": [29.8938, 17.1548, 34.9829, 16.3561, 55.6034, 65.1064],
    "params": 2.66, "flops": 10.3, "latency": 29.4,
}


def fmt(x):
    return f"{x:.2f}"


def main():
    d = Document(DOC)
    # table 6: A-D ablation summary
    t = d.tables[6]
    vals = [BASE, P2, None, None]
    labels = ["A：YOLO11n", "B：A+P2", "C：A+融合", "D：A+P2+融合"]
    for row, data, label in zip(t.rows[1:], vals, labels):
        row.cells[0].text = label
        if data is None:
            continue
        for cell, value in zip(row.cells[1:], [fmt(data["m50"]), fmt(data["m5095"]), "待定义", fmt(data["params"]), fmt(data["flops"]), fmt(data["latency"])]):
            cell.text = value
    # table 7: class AP/recall pairs
    t = d.tables[7]
    for row, data in zip(t.rows[1:], vals):
        if data is None:
            continue
        for cell, ap, rec in zip(row.cells[1:], data["ap"], data["rec"]):
            cell.text = f"{fmt(ap)}/{fmt(rec)}"
    for p in d.paragraphs:
        if p.text.startswith("目前尚未完成 DeepPCB 上的训练与测试"):
            p.text = "目前已完成一次 10 epoch、320 输入的 CPU 预备训练与测试。本稿将该结果用于核对流程，正式性能结论仍需增加训练轮次和随机种子后确认。"
        if "由于尚无训练结果" in p.text:
            p.text = p.text.replace("由于尚无训练结果，本稿将各实验项目明确列为待实测。", "目前仅完成一次预备训练；本稿将未运行的融合消融项目列为待实测。")
        if "本章规定训练、测试与统计方法；目前尚无 DeepPCB 实测结果" in p.text:
            p.text = p.text.replace("本章规定训练、测试与统计方法；目前尚无 DeepPCB 实测结果。实验执行前需记录", "本章规定训练、测试与统计方法；已完成一次 10 epoch、320 输入的 CPU 预备实测。该结果用于验证流程，正式论文仍需增加训练轮次和随机种子。实验执行中记录")
        if "各项性能数据、统计检验与最终结论均标为待实测" in p.text:
            p.text = p.text.replace("各项性能数据、统计检验与最终结论均标为待实测", "表5-2和表5-3已填入本次预备实测数据；统计检验与最终结论仍待正式实验")
        if "所有百分数与延迟表格在真实运行后填写" in p.text:
            p.text = p.text.replace("所有百分数与延迟表格在真实运行后填写", "表5-2和表5-3记录本次预备实测值，正式结果需在扩展实验后更新")
        if "表5-2 A—D 消融结果（所有数值待实测）" in p.text:
            p.text = p.text.replace("表5-2 A—D 消融结果（所有数值待实测）", "表5-2 A—D 消融结果（A、B为本次10 epoch预备实测）")
        if "表5-3 各类 AP 与召回率（待实测，单位：%）" in p.text:
            p.text = p.text.replace("表5-3 各类 AP 与召回率（待实测，单位：%）", "表5-3 各类 AP 与召回率（A、B为本次10 epoch预备实测，单位：%）")
        if "目前尚未执行训练与测试，因此不能断言" in p.text:
            p.text = p.text.replace("目前尚未执行训练与测试，因此不能断言", "目前仅完成一次 10 epoch、320 输入的预备训练与测试，样本划分和指标可复核，但训练轮次不足，因此不能断言")
    d.core_properties.subject = "DeepPCB YOLO11 实测预备结果；10 epoch、320 输入、CPU；仍需扩展正式实验"
    d.save(DOC)
    print(DOC)


if __name__ == "__main__":
    main()
